"""原子保存 JSON 與兩張工作表的檢驗報告。"""
from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime
from pathlib import Path, PureWindowsPath
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LABELS = {"PASS": "合格", "FAIL": "不合格", "UNKNOWN": "無法判定", "ERROR": "無法判定",
          "NO_DATA": "無法判定", "NO_SPEC": "無規格"}
JUDGMENT_STYLES = {"PASS": ("C6EFCE", "006100"), "FAIL": ("FFC7CE", "9C0006"),
                   "UNKNOWN": ("FFE699", "9C5700"), "ERROR": ("FFE699", "9C5700"),
                   "NO_DATA": ("FFE699", "9C5700"), "NO_SPEC": ("E7E6E6", "595959")}
SUMMARY_HEADERS = ["檢驗結果", "產品類別", "型號", "檔名", "不合格／無法判定原因", "使用的 spec", "檢驗時間", "處理耗時（秒）", "檔案連結"]
DETAIL_HEADERS = ["檔案整體結果", "產品類別", "型號", "檔名", "資料類型", "檢驗項目", "項目判定", "規格下限", "規格上限", "實測最小值", "實測最大值", "Marker 值", "單位", "備註"]


def report_link(path: str | Path) -> Path | PureWindowsPath:
    """Map report links only; never use the network path for file operations."""
    local_root = PureWindowsPath(os.getenv("REPORT_LOCAL_ROOT", "").strip() or r"C:\public")
    share_root = PureWindowsPath(os.getenv("REPORT_SHARE_ROOT", "").strip() or r"\\192.168.1.101\public")
    try:
        relative = PureWindowsPath(str(path)).relative_to(local_root)
    except ValueError:
        return Path(path).absolute()
    return share_root / relative


def write_json_result(output_dir: Path, source_name: str, result: dict[str, Any]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{source_name}.json"
    fd, temporary = tempfile.mkstemp(dir=output_dir, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(result, stream, ensure_ascii=False, indent=2)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return path


def _append(ws: Any, values: list[Any], judgment: str) -> None:
    ws.append([ILLEGAL_CHARACTERS_RE.sub("", value)[:32767] if isinstance(value, str) else value for value in values])
    fill, color = JUDGMENT_STYLES.get(judgment, JUDGMENT_STYLES["UNKNOWN"])
    for cell in ws[ws.max_row]:
        # Filenames, error text and model responses are always literal text.
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Microsoft JhengHei", size=11, color=color)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _setup(ws: Any, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Microsoft JhengHei", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "E2"
    ws.sheet_view.showGridLines = False


def write_excel_report(results: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "檔案總表"
    detail = workbook.create_sheet("檢驗明細")
    _setup(summary, SUMMARY_HEADERS, [16, 16, 23, 27, 65, 55, 23, 19, 55])
    _setup(detail, DETAIL_HEADERS, [18, 16, 23, 27, 16, 24, 16, 16, 16, 18, 18, 16, 12, 65])
    for result in results:
        overall = result.get("overall", "UNKNOWN")
        source = Path(result["source_file"])
        identity = [LABELS.get(overall, "無法判定"), result.get("category", ""), result.get("product_model", ""), source.name]
        inspected = result.get("inspected_at")
        if inspected:
            inspected = datetime.fromisoformat(inspected).replace(tzinfo=None)
        reason = "；".join(result.get("reasons") or []) or result.get("error", "")
        if result.get("archive_error"):
            reason = "；".join(filter(None, [reason, result["archive_error"]]))
        link = report_link(result.get("file_link") or source)
        spec_file = str(report_link(result["spec_file"])) if result.get("spec_file") else ""
        _append(summary, identity + [reason, spec_file, inspected,
                                   result.get("elapsed_seconds"), str(link)], overall)
        summary.cell(summary.max_row, 7).number_format = "yyyy-mm-dd hh:mm:ss"
        summary.cell(summary.max_row, 8).number_format = "0.00"
        summary.cell(summary.max_row, 9).hyperlink = link.as_uri()
        summary.cell(summary.max_row, 9).font = Font(name="Microsoft JhengHei", size=11, color="0563C1", underline="single")
        count_before = detail.max_row
        frequency = result.get("frequency") or {}
        if frequency:
            for key, label in (("start", "掃頻起點"), ("stop", "掃頻終點")):
                judgment = frequency.get(f"{key}_judgment", frequency.get("judgment", "NO_SPEC"))
                required = frequency.get(f"required_{key}_ghz")
                _append(detail, identity + ["掃頻範圍", label, LABELS.get(judgment, "無法判定"),
                        required if key == "stop" else None, required if key == "start" else None,
                        None, None, frequency.get(f"{key}_ghz"), "GHz", frequency.get("note", "")], judgment)
        for row_type, entries in (("整條曲線", result.get("trace_extrema") or []), ("Marker", result.get("measurements") or [])):
            for item in entries:
                judgment = item.get("judgment", "UNKNOWN")
                _append(detail, identity + [row_type, item["name"], LABELS.get(judgment, "無法判定"),
                        item.get("spec_min"), item.get("spec_max"), item.get("min"), item.get("max"),
                        item.get("value"), item.get("unit"), item.get("note", "")], judgment)
        if detail.max_row == count_before:
            _append(detail, identity + ["處理異常", "檔案檢驗", "無法判定", None, None, None, None, None, None, reason], "UNKNOWN")
        # Column A is the file result even when a particular item passes.
        fill, color = JUDGMENT_STYLES.get(overall, JUDGMENT_STYLES["UNKNOWN"])
        for row_index in range(count_before + 1, detail.max_row + 1):
            cell = detail.cell(row_index, 1)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name="Microsoft JhengHei", size=11, color=color)
    for ws in (summary, detail):
        ws.auto_filter.ref = ws.dimensions
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
